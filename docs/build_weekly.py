"""
Build the weekly reports for one week, one file per team member.

The previous week's report is used as the template so the university's layout,
fonts and signature block carry over untouched; only the dates, the two bullet
lists and the references are rewritten.

References are held per member per week and checked against every earlier week
before a file is written. The guide's review found the same three sources
repeated in Weeks 1 and 2 and again in Weeks 3 and 4, so a repeat is an error
here rather than something to be noticed after submission.

    python docs/build_weekly.py
"""
import copy
import os
import re
import subprocess

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

HERE = os.path.dirname(os.path.abspath(__file__))
DOWNLOADS = os.path.expanduser("~/Downloads")
SOFFICE = "/Applications/LibreOffice.app/Contents/MacOS/soffice"

WEEK = 12
FROM_DATE, TO_DATE = "20-09-2026", "25-09-2026"   # Sunday to Friday
NEXT_FROM, NEXT_TO = "27-09-2026", "02-10-2026"   # Sunday to Friday

REPORTS = {
    "24DCE052": {
        "work": [
            "Closed the scoring endpoint, which had been open to anyone who could reach the "
            "server. It is not a read: every call stores a transaction, extends the sender's "
            "payment history and adds entries to the network, so an open endpoint let a "
            "stranger corrupt the very history every later decision is judged against. It "
            "now requires the shared key, as the other writing endpoints already did.",
            "Wrote a test that walks the application's own list of endpoints and requires "
            "every one that writes to ask for the key, rather than testing the guard by "
            "itself, because the failure that actually happens is an endpoint that forgets "
            "to ask. Writing it found a second open one: account creation, which accepts the "
            "risk tier and the average spend that the rule engine scores against. Both are "
            "closed now, and the test fails if a new writing endpoint is added without the "
            "guard.",
            "Made the key fail closed. It had a working default written in the source and "
            "the same value again in the dashboard, so the guard could be passed by reading "
            "the project. There is no default now: development keeps a known key for "
            "convenience, and every other setting must supply its own, with the published "
            "value refused as firmly as no value at all. What this cannot fix is written "
            "down rather than implied, since the dashboard is a browser page and carries "
            "its copy of the key in the download: the guard keeps the writing endpoints "
            "closed to the open internet, not to a determined user.",
            "Stored what each layer contributed to a decision. The reply had always carried "
            "the three scores separately, but only their total was kept, so no past decision "
            "could be traced back to the layer that caused it. Each is saved beside the "
            "total now, and left empty rather than zero for a layer that did not run, so a "
            "layer that failed cannot be mistaken for a layer that found nothing.",
            "Gave a human override a record. Approving or blocking a transaction by hand "
            "used to overwrite the decision and leave nothing behind, so nobody could say "
            "afterwards that a block had been released, or by whom. Each override now "
            "appends the decision it replaced, the time, the person the caller declares and "
            "a reason if one is given; an override with nothing stated is still recorded and "
            "shows that nothing was stated.",
            "Finished the eight hundred and ninety-six transactions whose written "
            "explanation had never arrived. The language-model step runs after the answer is "
            "returned, inside the server, so a restart or a timeout left the record "
            "incomplete and nothing retried it. A script now completes them through the same "
            "decision path the live system uses and stamps each with the time it was "
            "finished, so an explanation written weeks later is not read as belonging to the "
            "moment of the decision. All eight hundred and ninety-six completed, none "
            "failed.",
        ],
        "plans": [
            "Fill the payment history from the stored transactions, so an account is not "
            "treated as new the first time it pays after a restart. Carried from last week "
            "and still open.",
            "Run the public-dataset replay once the file is downloaded and report each "
            "rule's true and false alarms on it. Carried from last week; the replay is "
            "written but the dataset has not been fetched.",
            "Submit the report and the requirement specification with the member "
            "responsible named on each chapter, and rehearse the demonstration from a "
            "cleanly started system.",
        ],
        "references": [
            "J. H. Saltzer and M. D. Schroeder, “The protection of information in computer "
            "systems,” Proceedings of the IEEE, vol. 63, no. 9, 1975.",
            "Open Worldwide Application Security Project, “OWASP API security top 10 — "
            "API2:2023 broken authentication,” OWASP, 2023. [Online]. Available: owasp.org",
            "K. Kent and M. Souppaya, “Guide to computer security log management,” NIST "
            "Special Publication 800-92, National Institute of Standards and Technology, "
            "2006.",
            "Reserve Bank of India, “Master direction on digital payment security "
            "controls,” RBI, 2021. [Online]. Available: rbi.org.in",
        ],
    },
    "24DCE040": {
        "work": [
            "The network page had been showing counts and no network: how many accounts, "
            "devices and addresses exist, with nothing drawn. Added an endpoint that returns "
            "the accounts worth looking at, the known-fraudulent ones first and then the "
            "most connected, so the page never opens on an account with nothing around it.",
            "The page now draws one hop around the chosen account — the accounts it paid, "
            "and the devices and addresses it shares — with a fraudulent neighbour and the "
            "link to it marked. Choosing another account redraws it. On one of the seeded "
            "ring accounts it shows two fraudulent neighbours, a shared device and a shared "
            "address together, so the ring is visible in a picture rather than described in "
            "a sentence.",
            "Ring detection was measured again in this week's fresh evaluation and flagged "
            "thirty-six of the thirty-six ring transactions. Recording again that this is "
            "the whole layer's figure and not the cycle check's own: twenty-one came from "
            "the cycle check and the rest from the receiver's closeness to an account "
            "already known to be fraudulent.",
            "The correction to the cycle query, designed last week, has not been written. "
            "It is left for the coming week deliberately rather than rushed into the week "
            "the measurements were taken, because changing the query and re-measuring in "
            "the same week would leave no way to tell which figure belongs to which "
            "version.",
        ],
        "plans": [
            "Write the correction so the query also matches the transfer that closes a "
            "ring, and measure the cycle check on its own afterwards rather than reporting "
            "the layer's total.",
            "Load the public dataset's accounts and transfers into the network, so the "
            "cycle and proximity checks are tested on payments we did not generate.",
            "Prepare the graph layer's part of the final demonstration, including the "
            "account whose neighbourhood shows a ring most clearly.",
        ],
        "references": [
            "B. Shneiderman, “The eyes have it: a task by data type taxonomy for "
            "information visualizations,” Proc. IEEE Symposium on Visual Languages, 1996.",
            "F. van Ham and A. Perer, “Search, show context, expand on demand: supporting "
            "large graph exploration with degree-of-interest,” IEEE Transactions on "
            "Visualization and Computer Graphics, vol. 15, no. 6, 2009.",
            "I. Herman, G. Melançon and M. S. Marshall, “Graph visualization and navigation "
            "in information visualization: a survey,” IEEE Transactions on Visualization "
            "and Computer Graphics, vol. 6, no. 1, 2000.",
            "I. Robinson, J. Webber and E. Eifrem, “Graph Databases: New Opportunities for "
            "Connected Data,” 2nd ed., O'Reilly Media, 2015.",
        ],
    },
    "24DCE029": {
        "work": [
            "Split the console into five pages behind the sidebar — overview, review queue, "
            "signals, model accuracy and the network — because it had grown to eleven "
            "panels on one scrolling page and the sidebar scrolled to them rather than "
            "navigating. Each page has its own address, so the back button, a reload and a "
            "bookmark all return to the page that was open. The figures are still fetched "
            "once for the whole console, so moving between pages costs no further requests.",
            "The accuracy page now reports four measures instead of three: precision, "
            "recall, accuracy and the F1 score. Accuracy is computed from the decisions "
            "that matched their label, and written beside it is the reason it must be read "
            "next to the others rather than instead of them — it counts approvals too, so "
            "on a set that is mostly legitimate it stays high whatever the detector does.",
            "Ran the evaluation again against all thirty rules, since the figures on display "
            "had been measured when there were five. Recall rose from 86.5 to 100 per cent: "
            "the seven frauds that used to be approved are all caught now. Precision moved "
            "from 97.8 to 96.3 per cent, F1 to 98.1 and accuracy to 99.0.",
            "Reported what that set cannot answer, rather than presenting the improvement "
            "alone. The precision figure fell because of one extra false alarm among a "
            "hundred and fifty-six legitimate payments, which is within measurement error "
            "and not evidence of anything. With only two false alarms in the set, the false "
            "alarm rate is too uncertain to carry to a real population: at a realistic fraud "
            "rate of one in a thousand the same detector's precision lies somewhere between "
            "about two and twenty-two per cent, and the single number it appears to be "
            "should not be quoted on its own.",
            "Added the decision history to the transaction panel, so an override and the "
            "decision it replaced are visible together, and marked any explanation that was "
            "written later than the decision it explains with the date it was written.",
        ],
        "plans": [
            "Measure on the public dataset, where a hundred and fifty-six legitimate "
            "payments become tens of thousands and the false alarm rate can be resolved "
            "well enough to quote.",
            "Continue the accessibility pass over the five pages, keyboard order first.",
            "Prepare the console for the final demonstration and check every page against a "
            "freshly started system rather than one that has been running all day.",
        ],
        "references": [
            "T. Saito and M. Rehmsmeier, “The precision-recall plot is more informative "
            "than the ROC plot when evaluating binary classifiers on imbalanced datasets,” "
            "PLOS ONE, vol. 10, no. 3, 2015.",
            "E. B. Wilson, “Probable inference, the law of succession, and statistical "
            "inference,” Journal of the American Statistical Association, vol. 22, no. 158, "
            "1927.",
            "L. D. Brown, T. T. Cai and A. DasGupta, “Interval estimation for a binomial "
            "proportion,” Statistical Science, vol. 16, no. 2, 2001.",
            "D. M. W. Powers, “Evaluation: from precision, recall and F-measure to ROC, "
            "informedness, markedness and correlation,” Journal of Machine Learning "
            "Technologies, vol. 2, no. 1, 2011.",
        ],
    },
}

def _set_text(paragraph, text):
    """Replace a paragraph's text, keeping the first run's formatting."""
    for run in paragraph.runs[1:]:
        run._r.getparent().remove(run._r)
    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


def _rewrite(paragraphs, lines):
    """Replace a run of paragraphs with `lines`, cloning the first for format."""
    first = paragraphs[0]
    for extra in paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    _set_text(first, lines[0])
    cursor = first._p
    for line in lines[1:]:
        clone = copy.deepcopy(first._p)
        cursor.addnext(clone)
        cursor = clone
        _set_text(Paragraph(clone, first._parent), line)


def _reference_paragraphs(doc):
    """The numbered paragraphs that follow the References heading."""
    out, seen_heading = [], False
    for p in doc.paragraphs:
        if p.text.strip().startswith("References"):
            seen_heading = True
            continue
        if seen_heading:
            if re.match(r"^\d+\.", p.text.strip()):
                out.append(p)
            elif out:
                break
    return out


def _key(reference):
    """Compare references loosely: lower case, letters and digits only."""
    return re.sub(r"[^a-z0-9]+", "", reference.lower())[:60]


def earlier_references(student_id):
    """Every reference this student has already cited, keyed to its week."""
    seen = {}
    for week in range(1, WEEK):
        docx = os.path.join(DOWNLOADS, f"{student_id}_Week{week}.docx")
        pdf = os.path.join(DOWNLOADS, f"{student_id}_Week{week}.pdf")
        if os.path.exists(docx):
            lines = [re.sub(r"^\s*\d+\.\s*", "", p.text)
                     for p in _reference_paragraphs(Document(docx))]
        elif os.path.exists(pdf):
            text = subprocess.run(["pdftotext", "-layout", pdf, "-"],
                                  check=True, capture_output=True, text=True).stdout
            lines = [re.sub(r"^\s*\d+\.\s*", "", ln)
                     for ln in text.splitlines() if re.match(r"^\s*\d+\.\s", ln)]
        else:
            continue
        for line in lines:
            seen.setdefault(_key(line), week)
    return seen


def _template_source(student_id):
    """The most recent earlier week's .docx still on disk, or None. A submitted
    week's .docx is often cleaned out of Downloads afterward (only the PDF is
    kept), so WEEK-1 isn't always there — walk back to whichever one is, and
    fall back to rebuilding the layout when every one of them is gone."""
    for week in range(WEEK - 1, 0, -1):
        path = os.path.join(DOWNLOADS, f"{student_id}_Week{week}.docx")
        if os.path.exists(path):
            return path
    return None


# ── Rebuilding the university layout ──────────────────────────────────────────
# Used when no earlier .docx survives. The two header logos were lifted out of a
# submitted Week 10 PDF (pdfimages) and live in docs/weekly-template/, so the
# layout can be rebuilt without one. Copying an earlier .docx is still preferred
# — it carries the exact original styles — so this runs only as a fallback.
TEMPLATE_DIR = os.path.join(HERE, "weekly-template")
TIMES, TITLE_FONT = "Times New Roman", "Calibri"


def _run(paragraph, text, size=11, bold=False, font=TIMES):
    r = paragraph.add_run(text)
    r.font.name, r.font.size, r.bold = font, Pt(size), bold
    r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    return r


def _line(doc, text="", size=11, bold=False, align=None, after=6, font=TIMES, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(space_before)
    if align is not None:
        p.alignment = align
    if text:
        _run(p, text, size, bold, font)
    return p


def _boxed(doc, lines):
    """One bordered cell holding the bullet list, as in the original form."""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.cell(0, 0)
    cell.text = ""
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _run(p, f"\u2022  {line}", 11)
    return t


def build_from_scratch(student_id, spec):
    """Rebuild the weekly-report layout: header with both logos, the details
    table, the two bordered lists, references and the signature line."""
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Inches(1))
    section.top_margin = section.bottom_margin = Inches(0.8)

    head = doc.add_table(rows=1, cols=3)
    head.autofit = False
    # Default cell padding would push the right-hand logo past the margin.
    margins = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), "0")
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    head._tbl.tblPr.append(margins)
    widths = (0.8, 4.9, 0.8)     # 6.5 in of usable width between 1 in margins
    for column, w in zip(head._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        column.set(qn("w:w"), str(int(w * 1440)))
    for i, w in enumerate(widths):
        head.rows[0].cells[i].width = Inches(w)
    head.cell(0, 0).paragraphs[0].add_run().add_picture(
        os.path.join(TEMPLATE_DIR, "charusat.png"), width=Inches(0.6))
    middle = head.cell(0, 1).paragraphs[0]
    middle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(middle, "CHAROTAR UNIVERSITY OF SCIENCE & TECHNOLOGY", 11.5, True)
    second = head.cell(0, 1).add_paragraph()
    second.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(second, "DEVANG PATEL INSTITUTE OF ADVANCE TECHNOLOGY AND RESEARCH", 8.5, True)
    right = head.cell(0, 2).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    right.add_run().add_picture(os.path.join(TEMPLATE_DIR, "depstar.png"), width=Inches(0.7))

    _line(doc, "WEEKLY REPORT", 16, True, WD_ALIGN_PARAGRAPH.CENTER, after=10,
          font=TITLE_FONT, space_before=16)

    details = doc.add_table(rows=3, cols=2)
    details.style = "Table Grid"
    rows = [(f"Project  ID: PRJ_CE_5_2026_7", f"Student ID: {student_id}"),
            (f"From Date: {FROM_DATE}", f"To Date: {TO_DATE}"),
            ("Semester: 5 th", "Internship ID:")]
    for r, (left_text, right_text) in enumerate(rows):
        for c, text in ((0, left_text), (1, right_text)):
            cell = details.cell(r, c)
            cell.text = ""
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)
            _run(cell.paragraphs[0], text, 11, True)
    for column, w in zip(details._tbl.tblGrid.findall(qn("w:gridCol")), (3.2, 2.8)):
        column.set(qn("w:w"), str(int(w * 1440)))
    for row in details.rows:
        row.cells[0].width, row.cells[1].width = Inches(3.2), Inches(2.8)

    _line(doc, f"Work done from Date: {FROM_DATE} to {TO_DATE}", 12, True, after=8, space_before=14)
    _boxed(doc, spec["work"])
    _line(doc, f"Plans for next week: Date: {NEXT_FROM} to {NEXT_TO}", 12, True,
          after=8, space_before=14)
    _boxed(doc, spec["plans"])

    _line(doc, "References:", 12, True, after=6, space_before=14)
    for i, reference in enumerate(spec["references"], 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Inches(0.3)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _run(p, f"{i}.  {reference}", 11)

    _line(doc, after=0, space_before=36)
    signature = doc.add_table(rows=1, cols=2)
    for i, text in enumerate(("Signature of Student", "Signature of Mentor")):
        cell = signature.cell(0, i)
        cell.text = ""
        cell.width = Inches(3.0)
        _run(cell.paragraphs[0], text, 12, True)

    out = os.path.join(DOWNLOADS, f"{student_id}_Week{WEEK}.docx")
    doc.save(out)
    return out


def build(student_id, spec):
    source = _template_source(student_id)
    if source is None:
        return build_from_scratch(student_id, spec)
    doc = Document(source)

    dates = doc.tables[1]
    # The template's first column is a shade too narrow for the project ID, so
    # "PRJ_CE_5_2026_7" wraps onto a second line in some renderers. Widen it.
    for column, width in zip(dates._tbl.tblGrid.findall(qn("w:gridCol")), (3.2, 2.8)):
        column.set(qn("w:w"), str(int(width * 1440)))
    for row in dates.rows:
        row.cells[0].width = Inches(3.2)
        row.cells[1].width = Inches(2.8)
    _set_text(dates.cell(1, 0).paragraphs[0], f"From Date: {FROM_DATE}")
    _set_text(dates.cell(1, 1).paragraphs[0], f"To Date: {TO_DATE}")

    for p in doc.paragraphs:
        if p.text.startswith("Work done from Date:"):
            _set_text(p, f"Work done from Date: {FROM_DATE} to {TO_DATE}")
        elif p.text.startswith("Plans for next week:"):
            _set_text(p, f"Plans for next week: Date: {NEXT_FROM} to {NEXT_TO}")

    _rewrite(doc.tables[2].cell(0, 0).paragraphs,
             [f"•  {line}" for line in spec["work"]])
    _rewrite(doc.tables[3].cell(0, 0).paragraphs,
             [f"•  {line}" for line in spec["plans"]])
    _rewrite(_reference_paragraphs(doc),
             [f"{i}.  {r}" for i, r in enumerate(spec["references"], 1)])

    out = os.path.join(DOWNLOADS, f"{student_id}_Week{WEEK}.docx")
    doc.save(out)
    return out


if __name__ == "__main__":
    for student_id, spec in REPORTS.items():
        already = earlier_references(student_id)
        repeats = [r for r in spec["references"] if _key(r) in already]
        if repeats:
            raise SystemExit(f"{student_id}: already cited in an earlier week — {repeats}")
        out = build(student_id, spec)
        subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf",
                        "--outdir", DOWNLOADS, out], check=True, capture_output=True)
        print(f"{student_id}: {os.path.basename(out)} + .pdf — "
              f"{len(spec['references'])} references, none of the "
              f"{len(already)} cited in Weeks 1-{WEEK - 1}")
