"""
Build one contribution document per team member.

Each document answers three questions for a reader who has not seen the code:
what this member owned, what each of those components actually does, and what
they did in each week of the semester.

The timeline is taken from the weekly reports the member submitted, so it
records what was written at the time rather than what anyone remembers now.
The component descriptions are written against the code as it stands.

    python docs/build_contributions.py
"""
import os
import subprocess

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
DOWNLOADS = os.path.expanduser("~/Downloads")
SOFFICE = "/Applications/LibreOffice.app/Contents/MacOS/soffice"

FONT = "Times New Roman"
BODY, SUB, TITLE = 12, 13, 16
INK = RGBColor(0, 0, 0)

PROJECT = "PRJ_CE_5_2026_7"
SUBTITLE = "Indus11 — AI Financial Risk and Fraud Decision Engine"


# ── document plumbing ────────────────────────────────────────────────────────
def new_document():
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(1)
        section.top_margin = section.bottom_margin = Inches(1)
    return doc


def para(doc, text, size=BODY, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=8):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text)
    run.font.size, run.bold, run.font.name, run.font.color.rgb = Pt(size), bold, FONT, INK
    return p


def heading(doc, text, size=SUB):
    p = para(doc, text, size, True, WD_ALIGN_PARAGRAPH.LEFT, after=6)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.keep_with_next = True
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    run.font.size, run.font.name, run.font.color.rgb = Pt(BODY), FONT, INK
    return p


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    # Word honours a column width only when autofit is off and the width is set
    # on the grid as well as on every cell; setting cell.width alone is ignored.
    t.autofit = False
    for column, width in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        column.set(qn("w:w"), str(int(width * 1440)))
    for cell, text, width in zip(t.rows[0].cells, headers, widths):
        cell.width = Inches(width)
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold, run.font.size, run.font.name = True, Pt(10.5), FONT
        cell.paragraphs[0].paragraph_format.line_spacing = 1.0
    for row in rows:
        cells = t.add_row().cells
        for cell, text, width in zip(cells, row, widths):
            cell.width = Inches(width)
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(text))
            run.font.size, run.font.name = Pt(10.5), FONT
            cell.paragraphs[0].paragraph_format.line_spacing = 1.0
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return t


# ── the three members ────────────────────────────────────────────────────────
OM = {
    "id": "24DCE052",
    "name": "Joshi Om",
    "role": "Gateway, data storage, rule engine and project infrastructure",
    "summary":
        "Om built the part of the system every transaction passes through first and last. A "
        "payment arrives at his gateway, is checked for shape and sanity, is scored by his "
        "rule engine against thirty separate fraud indicators, and is written to the "
        "database by his persistence code. He also owns the machinery around the project: "
        "the containers, the one-command local stack, the automated tests, and the "
        "continuous-integration workflow that runs them on every push.",
    "components": [
        ("The gateway, Layer 1", "app/api/routes/transactions.py",
         "This is the front door. When a payment system sends a transaction, the gateway "
         "receives it. Before any scoring happens it checks that the request makes sense: "
         "that the amount is positive, that the currency code is three letters, that the "
         "merchant category is one the system knows, and that no free-text field is long "
         "enough to be an attack. A request failing any of these is refused with a message "
         "saying why, so a malformed payment never reaches the detection layers.",
         "The gateway then loads the profiles of the sender and the receiver. It looks in "
         "Redis first, because that answers in well under a millisecond, and falls back to "
         "MongoDB when the cache has no copy. An account nobody has seen before is given a "
         "default profile marked elevated risk, since an unknown party is not the same thing "
         "as a safe one."),
        ("The rule engine, Layer 2", "app/services/rule_engine.py",
         "The rule engine is the system's fastest and most certain opinion. It awards up to "
         "forty points out of a hundred and decides in a few milliseconds, because every "
         "check it performs is a direct lookup rather than a model prediction. Five of its "
         "rules were there from the start: a blacklisted party, more than five payments in "
         "ten minutes, an amount more than three times the account's own average, a "
         "high-risk merchant category, and the sender's risk tier.",
         "Twenty-five further rules arrived after the mid-semester review, drawn from the "
         "red-flag indicators the Financial Intelligence Unit of India and the Reserve Bank "
         "publish for suspicious-transaction reporting, from the international "
         "money-laundering typologies, and from card-network fraud patterns. They cover the "
         "shapes a human investigator looks for: an amount kept just under the ten-lakh "
         "reporting limit, several smaller payments that together cross it, an account "
         "dormant for six months suddenly active, money received and passed straight on "
         "again, payments to many new beneficiaries in one day, a payment from a device or a "
         "place the account has never used, two payments too far apart in distance to be the "
         "same person, and a payment note written in the language a scam uses.",
         "However many rules fire, the layer's total is capped at forty. That cap is "
         "deliberate. It means the rule engine alone can never push a transaction into the "
         "block band, so no single kind of evidence stops a payment without corroboration "
         "from another layer."),
        ("Account history in Redis", "app/core/redis_client.py",
         "Most anomaly rules ask about the past rather than the present. Has this sender "
         "paid this person before? How many different people have they paid today? How much "
         "money came in during the last day? Which devices has this account used? Redis "
         "holds the answers in sorted sets, which keep entries in time order and let old "
         "ones age out exactly instead of being cleared in bulk.",
         "All of it is read in a single round trip, with the reads queued before the writes. "
         "That ordering matters. It guarantees every answer describes the account as it was "
         "immediately before this payment rather than after it. Without it, a rule asking "
         "whether the sender has used this device before would always answer yes, because "
         "the current payment would already have recorded the device."),
        ("Storage and the decision record", "app/models/transaction.py",
         "Every analysed transaction becomes a document in MongoDB through the Beanie object "
         "document mapper. The record holds the composite score, the decision, the written "
         "explanation, and which layers failed if any did.",
         "Two late additions made a past decision reconstructable. The score each layer "
         "contributed is now stored beside the total, so a decision can be traced to the "
         "layer responsible for it, and a layer that did not run is stored as empty rather "
         "than zero so a failure cannot be mistaken for a clean result. A human override now "
         "appends a record of what it replaced, when, who the caller says they are and why, "
         "because the decision field itself is overwritten and the log is the only place the "
         "previous value survives."),
        ("Keeping the write routes closed", "app/core/security.py",
         "Five routes change stored state, and all of them require a shared key in the "
         "X-API-Key header, compared in constant time so the comparison cannot leak the key "
         "through its own timing. The scoring endpoint is one of them. It looks like a read, "
         "but every call stores a transaction, extends the sender's history and adds entries "
         "to the graph, so leaving it open would let a stranger corrupt the very history "
         "every later decision is judged against.",
         "What a shared key cannot do is written down rather than implied. It identifies "
         "nobody, and the dashboard is a browser page that has to carry its copy in the "
         "download, so the guard closes the write routes to the open internet and not to a "
         "determined user. Per-user authentication is recorded as out of scope for the "
         "semester."),
        ("Measuring the system honestly", "scripts/evaluate.py",
         "The evaluation harness replays a labelled set of two hundred and eight "
         "transactions through the running API and reports precision, recall, F1 and a full "
         "confusion matrix, then sweeps the decision thresholds to show what other bands "
         "would have produced.",
         "Its most useful feature is the one that makes the numbers look worse. The test set "
         "is a quarter fraud while a real payment feed is nearer one in a thousand, so the "
         "harness also reports what the same detector's precision would be at that "
         "prevalence. Both figures are printed, and both appear in the report."),
        ("The local stack, the tests and continuous integration", "scripts/run_local.sh",
         "One command brings up the three databases, the API and the dashboard together, "
         "which is what makes the system demonstrable on a laptop rather than only on the "
         "machine it was written on. The same stack is defined for Docker Compose.",
         "One hundred and thirteen automated tests run on every push through GitHub Actions. "
         "None needs a database or a network connection, which is why they can run in "
         "continuous integration at all. The most recent addition walks the application's "
         "own list of endpoints and fails the build if a route that writes has forgotten to "
         "ask for the key. Writing that test found an endpoint which had."),
    ],
    "timeline": [
        ("1", "06–12 Jul", "Worked out what the backend had to do and wrote the transaction API contract as Pydantic schemas. Set up the Python environment, the FastAPI skeleton and the folder structure, chose the backend stack, and prepared the Docker Compose configuration."),
        ("2", "13–19 Jul", "Designed the MongoDB document schema for accounts and transactions, settled the backend architecture and the analyse flow, specified the Redis cache keys and the velocity-window approach, and documented the rule-engine design."),
        ("3", "21–27 Jul", "Completed the gateway and the rule engine: validation, profile loading with a cache fallback, rate limiting, velocity checks on Redis sorted sets, and the amount, blacklist and merchant rules. Built the MongoDB data layer and a seed script for twenty test accounts, containerised the stack, and set up the repository with continuous integration."),
        ("4", "28 Jul – 3 Aug", "Built the first accuracy-evaluation harness and re-tuned the review and block thresholds from its output instead of from assumption. Hardened the API with strict payload validation and revised rate limits after load testing."),
        ("5", "4–10 Aug", "Wrote the first three chapters of the specification, the thirty-two functional requirements as a traceable table, and the database design chapter. Checked every schema table against the running code and corrected four errors, including an index claimed on a field that is not indexed. Automated the document build."),
        ("6", "9–14 Aug", "Presented at Review 1. The industry expert's central objection was that the measured fourteen-second pipeline put the language model inside the decision path, which no payment decision can afford. Agreed the fix with the team: return the deterministic decision inside a 500 ms budget and attach the explanation afterwards."),
        ("7", "16–21 Aug", "Implemented that split. The rule engine and graph analyzer now decide and persist on their own, and the language-model call moved to a background task that updates the same record when it finishes. Added a pending flag end to end so a caller can tell whether the explanation has landed."),
        ("8", "23–28 Aug", "Load-tested the endpoint now that it returns before the model call: the median response at twenty concurrent requests fell from about 14.6 seconds to under 400 ms. Added a server-sent-events endpoint so the dashboard learns of an update without polling."),
        ("9", "30 Aug – 4 Sep", "Traced a fault that only appears under sustained load. The background model call had no deadline, so four hung requests held every concurrency permit and silently stopped the whole background pipeline. Confirmed the diagnosis by measurement rather than assumption, then gave the call a 180-second deadline. Also corrected the harness, which had been recording the immediate response and omitting the model's contribution entirely."),
        ("10", "6–11 Sep", "Found that the language model had never been shown what the other layers detected, so its explanation was written blind to the signals printed beside it. The flags are now passed in as confirmed findings. Added precision at realistic prevalence to the evaluation, and measured each layer's contribution to recall: no layer detects any fraud on its own."),
        ("11", "13–18 Sep", "Catalogued fifty banking anomalies from the published red-flag indicators and implemented twenty-five of them. Kept the fifteen that cannot be detected in the report, each with the data it would need, rather than dropping them quietly. Found a fault in the pass-through rule by replaying scenarios through the running system."),
        ("12", "20–25 Sep", "Closed the scoring endpoint and, through a new test over the application's own route list, found account creation open as well. Made the key fail closed outside development. Stored each layer's score, gave overrides an audit record, and finished the eight hundred and ninety-six transactions whose explanation had never arrived."),
    ],
    "open": [
        "The account history in Redis begins at whatever moment the system was started, so an established account is treated as new the first time it pays after a restart. Filling it from the stored transactions is the fix, and it has not been done.",
        "The replay that scores the public PaySim dataset is written but has not been run, because the dataset has not been downloaded. Every accuracy figure in the project therefore still rests on data the team generated itself.",
        "The language-model layer runs inside the API process with no queue behind it. The drain script finishes what a restart interrupts, but that is a repair rather than a design.",
    ],
}

KRISH = {
    "id": "24DCE040",
    "name": "Krish Gajera",
    "role": "Graph database and relationship detection",
    "summary":
        "Krish owns the part of the system that looks at how accounts relate to one another "
        "rather than at any single payment. This catches the fraud the rule engine cannot "
        "see by design: a ring of accounts passing money around a loop, several "
        "unrelated-looking accounts operated from one device, or a receiver sitting close to "
        "an account already known to be fraudulent.",
    "components": [
        ("The transaction graph in Neo4j", "scripts/seed_neo4j.py",
         "A graph database stores things and the connections between them as first-class "
         "objects, which is what makes a question like “is there a path from this "
         "account back to itself” cheap to ask. The schema has three kinds of node, for "
         "accounts, devices and network addresses, joined by relationships recording that "
         "one account sent money to another, or that an account used a device or an address.",
         "Because no real transaction network was available, Krish wrote a generator that "
         "builds one: five hundred accounts, three money-mule rings of four, five and three "
         "accounts each cycling four times and losing between five and twelve per cent at "
         "every hop, two clusters sharing a device, one cluster of eight accounts sharing an "
         "address, and roughly nine hundred and forty ordinary transfers around them. It "
         "runs from a fixed seed, so every member of the team works with an identical graph."),
        ("What the graph layer looks for", "app/services/graph_analyzer.py",
         "The layer awards up to thirty points and checks five patterns. A shared device or "
         "a shared address means two or more accounts operating from one place. A circular "
         "flow means money returning to its sender through a chain of hops. A money-mule "
         "chain is the same shape with the amounts shrinking at each hop, which is what a "
         "fee being skimmed looks like. Cluster proximity means the receiver sits within two "
         "links of an account already labelled fraudulent.",
         "The cycle query carries two conditions that took measurement to get right. It only "
         "considers chains of two to four hops inside a seventy-two-hour window, and it "
         "requires the timestamps along the path to increase, because money moving around a "
         "ring must move forwards in time. Adding those conditions cut the pattern's false "
         "alarms from 44.9 per cent to 10.3 per cent and its query time from 260 ms to about "
         "20 ms."),
        ("Spreading a fraud label outwards", "app/services/graph_analyzer.py",
         "Knowing that one account is fraudulent says something about the accounts around "
         "it. The propagation routine materialises connections between accounts sharing a "
         "device or an address, then labels everything within two hops of a known fraud "
         "account as fraud-adjacent and records which cluster it belongs to. The graph layer "
         "can then score a transaction on how close its receiver sits to known trouble, "
         "which is the check that catches ring members the cycle query misses."),
        ("Graph endpoints for the dashboard", "app/api/routes/graph.py",
         "Three routes expose the graph. One returns the neighbourhood around a given "
         "account, one returns the entity counts, and one returns the accounts worth looking "
         "at, ranked with the known-fraudulent first and then by how connected each is, so "
         "the dashboard never opens on an account with nothing around it.",
         "That last route is what turned the dashboard's network page from a list of counts "
         "into an actual picture. Choosing a seeded ring account now draws two fraudulent "
         "neighbours, a shared device and a shared address together, so a reviewer sees the "
         "ring instead of reading a sentence about it."),
    ],
    "timeline": [
        ("1", "06–12 Jul", "Researched graph-database approaches to fraud-ring detection, installed Neo4j and got the asynchronous driver connecting, and studied the APATE paper and the Cypher patterns for shared-device and circular-flow detection."),
        ("2", "13–19 Jul", "Designed the graph schema of account, device and address nodes with transaction relationships, planned the structure of the synthetic dataset, and drafted the Cypher queries for shared-device and ring detection."),
        ("3", "21–27 Jul", "Implemented the schema with indexes created automatically at startup, wrote the dataset generator, and built the Cypher patterns for shared device and address, circular flows and fee-skimming cycles. Added the graph endpoints and the fraud-label propagation routine, and integrated the graph score into the composite decision."),
        ("4", "28 Jul – 3 Aug", "Validated ring detection against the planted rings and measured how many were recovered. Tuned the thresholds to reduce false alarms on ordinary repeat payments, and contributed ring-recall and cluster-precision figures to the shared evaluation."),
        ("5", "4–10 Aug", "Wrote the graph-layer section of the specification, describing the shared-device, circular-flow, fee-skimming and cluster-proximity checks and the Cypher patterns behind them."),
        ("6", "9–14 Aug", "Presented the graph layer at Review 1. Answered the question of which pattern produces the most false positives: shared-device detection, because several legitimate people can use one device. Began looking at what happens to the cycle query once the graph reaches a million records with no deletion path."),
        ("7", "16–21 Aug", "Added failure handling to the graph analyzer so a Neo4j outage degrades the layer instead of crashing the whole request, with a regression test that simulates the outage. This closed the gap the expert had raised about what the dashboard shows when the graph database goes down."),
        ("8", "23–28 Aug", "Ran the false-positive analysis across all four patterns on the seeded data. Benchmarked the cycle query as the graph grew and found its time rising from about 40 ms to 380 ms, traced it to an unindexed relationship property, added the index, and brought it back to about 60 ms."),
        ("9", "30 Aug – 4 Sep", "Confirmed by re-measurement that the change held: precision rose from 93.8 to 97.8 per cent and F1 from 0.900 to 0.918 with recall unchanged, the intended outcome since the change was meant to remove false alarms without losing detections. Ring detection reached every planted case, 36 of 36."),
        ("10", "6–11 Sep", "Separated what the graph layer as a whole detects from what the cycle check alone detects, and found the layer flagged all 36 ring transactions while the cycle check fired on only 21. Traced the cause: the query walks outward requiring times to rise, but the transfer closing a ring is always the newest, so it can never match."),
        ("11", "13–18 Sep", "Designed the correction and set out what it will cost to verify, since today's figure belongs to the layer as a whole and that distinction has to survive the next measurement. Planned how the public dataset would reach this layer, noting that its rows carry no device or address, so the shared-identity checks cannot be measured on it."),
        ("12", "20–25 Sep", "Gave the network page an actual network: an endpoint ranking the accounts worth exploring, and a drawing of one hop around the chosen account with fraudulent neighbours marked. Ring detection re-measured at 36 of 36 in the fresh evaluation."),
    ],
    "open": [
        "The correction to the cycle query is designed but not written. It was deliberately held back from the week the measurements were taken, because changing the query and re-measuring together would leave no way to tell which figure belongs to which version.",
        "A retention or archival policy for the transaction graph was discussed across several weeks and never settled. Without one the graph grows indefinitely and the cycle query slows with it.",
        "Fan-in and fan-out mule shapes were planned and not added, so their effect on ring recall is unknown.",
    ],
}

DRASHTI = {
    "id": "24DCE029",
    "name": "Drashti Dedaniya",
    "role": "Retrieval and language model, decision engine, the analyst dashboard, and the "
            "project's diagrams",
    "summary":
        "Drashti owns the two ends of the system a person actually sees: the written "
        "explanation that says why a transaction was held, and the screen the analyst works "
        "from. Between them sits the decision engine, which takes the three layers' scores, "
        "turns them into an approve, review or block decision, and refuses to publish an "
        "explanation that does not match the evidence. She also produced the data-flow and "
        "UML diagrams that the specification and the report are built around.",
    "components": [
        ("The knowledge base and retrieval", "app/services/fraud_kb.py",
         "Retrieval-augmented generation means giving a language model relevant reference "
         "material before asking it a question, so its answer is grounded in something "
         "rather than recalled from training. The reference material here is fifty-eight "
         "short documents describing fraud patterns, each typed by the kind of fraud it "
         "describes: synthetic identity, money mule, account takeover, wire fraud and others.",
         "They live in ChromaDB, a vector database, which finds documents by meaning rather "
         "than by keyword. A query built from the transaction's amount, merchant category, "
         "risk tier and device retrieves the three closest patterns, and those go into the "
         "prompt."),
        ("The language-model layer, Layer 4", "app/services/rag_pipeline.py",
         "The model returns a score out of thirty and a short written assessment. Two "
         "decisions shape how it behaves. Its temperature is fixed at zero, because at the "
         "default setting the same transaction scored five different ways in five runs, and "
         "Review 1 asked directly whether an identical transaction scores identically. It "
         "must, and now it does.",
         "The second came from a fault found in Week 10. The model had been given only the "
         "raw transaction, so it wrote its explanation without knowing what the rule and "
         "graph layers had detected, and its prose could contradict the flags printed beside "
         "it. The detected flags are now passed in as confirmed findings. On a transaction "
         "from a blacklisted sender using a device shared with eight accounts, the model had "
         "scored 0 and called the activity ordinary; it now scores 20 and cites both signals.",
         "The layer is also forgiving of its own failures. The reply is parsed by scanning "
         "for the first complete JSON object, because a local model will sometimes repeat an "
         "example before giving its own answer, and if the configured provider fails the "
         "pipeline retries against a local Ollama model so the system keeps working offline."),
        ("The decision engine, Layer 5", "app/services/decision_engine.py",
         "This adds the three layers' scores, caps the total at a hundred, and maps it to a "
         "decision: below forty approve, forty to sixty-nine review, seventy or above block. "
         "The bands are configuration rather than code, so they can be tuned without "
         "changing the logic.",
         "A layer that could not run is treated carefully. It contributes no points, but a "
         "transaction that would otherwise have been approved is sent for review instead, "
         "because a layer that never looked at the transaction is missing evidence rather "
         "than evidence of safety. A block already earned by the layers that did run is kept.",
         "The explanation guard is the part worth explaining to any reader. Language models "
         "produce fluent text whether or not it is true, so before an explanation is shown it "
         "is checked against the flags that actually fired. If the prose refers to none of "
         "them, only the flag list is shown, with a line saying the generated explanation was "
         "withheld. An analyst never receives a confident paragraph that contradicts the "
         "evidence behind the decision."),
        ("The analyst dashboard", "dashboard/src/",
         "The dashboard is a React application, now laid out as five pages behind a sidebar: "
         "an overview, the review queue, the signals breakdown, the model accuracy, and the "
         "graph network. Each page has its own address, so the back button, a reload and a "
         "bookmark all return to the page that was open.",
         "The review queue is the page an analyst actually works from. It lists recent review "
         "and block decisions with filters, a search and sortable columns, and shows the "
         "signals that fired as labelled chips rather than as clipped text. Selecting a row "
         "opens a panel with the transaction's details, its signals grouped by the layer that "
         "raised them, the written explanation, the history of any override, and a drawing of "
         "the accounts and devices one hop around the sender.",
         "Colour on this dashboard carries meaning and is never decoration. Only the three "
         "decisions appear in saturated colour, and those colours were chosen by measurement: "
         "the usual green and red fail the most common form of colour blindness, so an "
         "approval is shown in teal. Contrast was measured rather than judged by eye, which "
         "found the block colour at 3.93:1 and faint text at 3.24:1, both under the 4.5:1 "
         "minimum, and both were raised."),
        ("Reporting accuracy without overstating it", "dashboard/src/components/AccuracyPanel.jsx",
         "The accuracy page reports four measures, precision, recall, accuracy and the F1 "
         "score, with the confusion matrix beneath them. Written beside them is the reason "
         "accuracy has to be read next to the others rather than instead of them, since it "
         "counts approvals too and stays high on a set that is mostly legitimate whatever "
         "the detector does.",
         "The page also carries the figure the team could have left out. The test set is a "
         "quarter fraud; a real feed is nearer one in a thousand, and at that rate the same "
         "detector's precision falls to about seven per cent. That number sits on the screen "
         "next to the flattering one."),
        ("The diagrams", "docs/diagrams/",
         "The guide asked for correct notation, so the data-flow diagrams were redrawn in "
         "standard Gane\u2013Sarson form: a context diagram, a level 1 diagram with seven "
         "processes and five data stores, and a level 2 diagram expanding rule evaluation. "
         "The full UML set followed, covering use case, activity, sequence, class, state, "
         "component and deployment, along with the entity relationship diagram.",
         "She also measured what happened to each diagram once it was scaled to fit the page "
         "and found the class diagram's text had fallen to about five points. The class and "
         "activity diagrams were each split in two and two others reshaped, so nothing in "
         "print now falls below about eight points. The sources are committed alongside the "
         "images, so any diagram can be regenerated rather than redrawn by hand."),
    ],
    "timeline": [
        ("1", "13–19 Jul", "Designed the retrieval pipeline architecture, defined the decision engine's scoring bands, drew the dashboard wireframes, and structured the fraud-pattern knowledge base. (Submitted for the same dates as Week 2; see the note at the end.)"),
        ("2", "13–19 Jul", "As submitted, the same content and the same dates as Week 1."),
        ("3", "21–27 Jul", "Expanded the knowledge base to fifty-eight documents, improved the prompt with worked examples, and added the Ollama fallback. Implemented the retrieval pipeline and the decision engine, built the React dashboard with its charts and flagged-transactions table, added the transaction detail view with an override, and switched every amount to rupees."),
        ("4", "28 Jul – 3 Aug", "Added the accuracy panel showing precision, recall and the confusion matrix. Polished the dashboard with loading skeletons, empty states and clearer errors when the API is unavailable, and refined the prompt after reviewing explanations that did not match the triggered signals."),
        ("5", "4–10 Aug", "Made the flagged-transactions table interactive with filter chips, search and sortable columns. Added motion that switches itself off for a viewer who has reduced motion enabled. Built a static demo mode and verified it by stopping the API entirely and serving the build."),
        ("6", "9–14 Aug", "Brought the dashboard up against the live backend end to end for the first time. At Review 1, asked what an analyst sees if Neo4j goes down mid-session, and confirmed the case was not handled rather than glossing over it."),
        ("7", "16–21 Aug", "Added the explanation guard to the decision engine, so a written explanation referring to none of the triggered flags is withheld in favour of the flag summary. Fixed the accuracy panel's bare error state and began wiring the dashboard to the pending flag."),
        ("8", "23–28 Aug", "Finished the pending badge so a row clears itself once the explanation lands. Wrote an end-to-end test of the explanation guard against real model output rather than fixture text, which caught a bug where its flag matching was case-sensitive and silently fell back on mixed-case labels."),
        ("9", "30 Aug – 4 Sep", "Simplified the decision engine after review: two flags that always moved together became one, and the matching logic went from two functions to one. Removed state from the dashboard that could outlive the effect that created it."),
        ("10", "6–11 Sep", "Redesigned the dashboard around the rule that colour carries meaning. Measured contrast instead of judging it and raised two colours that failed the minimum. Fixed the decision chart breaking below 500 pixels. Found a flaw in the explanation guard: flags describing only the sender's risk tier contain nothing but generic words, so their explanations were always withheld, which affected 139 of the 156 legitimate transactions."),
        ("11", "13–18 Sep", "Rebuilt the analyst screen as a working console showing decisions per day, the decision mix, risk by merchant category, how often each signal fires and the score distribution, computed across every stored transaction by one aggregation endpoint rather than the twenty most recent. Corrected a fault that had made every timestamp five and a half hours early."),
        ("12", "20–25 Sep", "Split the console into five addressable pages. Added accuracy beside precision, recall and F1 with the reason it must not be read alone. Re-ran the evaluation against all thirty rules: recall rose from 86.5 to 100 per cent while precision moved from 97.8 to 96.3, which is one extra false alarm in 156 legitimate rows and within measurement error."),
    ],
    "open": [
        "The accessibility pass is unfinished. Contrast on the filters and the modal dialogs was still outstanding when the pass was last reported, and the keyboard-order review over the five new pages has not been done.",
        "The evaluation set contains only two false alarms among 156 legitimate transactions, too few to resolve the false-alarm rate that decides real-world alert quality. Measuring on a public dataset is the fix.",
        "The explanation guard matches a five-letter root of each distinctive word, which is a heuristic. It was tightened after the risk-tier flaw was found, but it has not been tested against a model deliberately trying to defeat it.",
        "Week 1 and Week 2 were submitted with the same content and the same dates, so one week of the record is missing.",
    ],
}


def build(member):
    doc = new_document()
    para(doc, "INDIVIDUAL CONTRIBUTION REPORT", TITLE, True, WD_ALIGN_PARAGRAPH.CENTER, after=4)
    para(doc, SUBTITLE, SUB, False, WD_ALIGN_PARAGRAPH.CENTER, after=14)
    table(doc,
          ["Student", "Student ID", "Project ID"],
          [[member["name"], member["id"], PROJECT]],
          [2.4, 2.0, 2.1])
    para(doc, member["role"], BODY, True, WD_ALIGN_PARAGRAPH.CENTER, after=14)

    heading(doc, "1.  What this work covers")
    para(doc, member["summary"])

    heading(doc, "2.  The components, and what each one does")
    for i, component in enumerate(member["components"], 1):
        name, path, *paragraphs = component
        heading(doc, f"2.{i}  {name}", BODY)
        para(doc, f"Where it lives: {path}", BODY - 1, False, WD_ALIGN_PARAGRAPH.LEFT, after=4)
        for text in paragraphs:
            para(doc, text)

    doc.add_page_break()
    heading(doc, "3.  Week by week")
    para(doc, "Taken from the weekly reports submitted at the time, so this records what was "
              "written then rather than what is remembered now.")
    table(doc, ["Week", "Dates", "Work done"], member["timeline"], [0.6, 1.1, 4.8])

    heading(doc, "4.  What is still open")
    para(doc, "Recorded here rather than left out, because a reader judging the work should "
              "be able to see where it stops.")
    for item in member["open"]:
        bullet(doc, item)

    out = os.path.join(DOWNLOADS, f"{member['id']}_Contribution.docx")
    doc.save(out)
    return out


if __name__ == "__main__":
    for member in (OM, KRISH, DRASHTI):
        out = build(member)
        subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf",
                        "--outdir", DOWNLOADS, out], check=True, capture_output=True)
        print(f"{member['id']}: {os.path.basename(out)} + .pdf  "
              f"({len(member['components'])} components, {len(member['timeline'])} weeks)")
