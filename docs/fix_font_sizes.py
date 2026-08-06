"""
Make font sizes consistent across equivalent elements in the review deck —
nothing else changes: no text, color, position, or layout edits.

Two genuine inconsistencies exist, found by surveying every run's font size:

1. Slide titles vary 24-30 pt depending on which script added the slide
   (24 pt: Process Flow + the six diagram slides; 26 pt: Literature Review
   and both Implementation slides; 30 pt: Outline; 28 pt: everything else).
   Normalized to 28 pt, the size the majority already use. The cover title
   on slide 1 is a distinct design element (a different placeholder role,
   not a "slide title" in this sense) and is left untouched.

2. The five "Two Content" bullet slides (Feedback, Problem Statement,
   Non-Functional Requirements, Progress Since the Last Review, Challenges)
   use two different scales: 18/16 pt on three of them, 15/13 pt on two.
   Normalized to 18/16, matching the majority. One stray heading on the
   NFR slide ("Software quality attributes") was already at 16 pt instead
   of the 18 pt its sibling headings use on the same slide — folded into
   the same fix since it's the same inconsistency.

Every other size in the deck (kicker labels, stat-card numbers, table
cells, code blocks, diagram callout boxes, the References list, the
Outline numbers) was already consistent within its own role and is left
exactly as it is.

    python docs/fix_font_sizes.py
"""
import os

from pptx import Presentation
from pptx.util import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
DECK = os.path.join(HERE, "Indus11-Review2.pptx")

TITLE_SIZE = Pt(28)
HEADER_SIZE = Pt(18)
BULLET_SIZE = Pt(16)
TWO_CONTENT_SLIDES = {3, 4, 6, 19, 20}   # 1-indexed

prs = Presentation(DECK)
slides = prs.slides

title_fixed = body_fixed = 0

for i, slide in enumerate(slides, 1):
    if i == 1:
        continue  # cover slide title is a different design role, untouched
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 0 and ph.has_text_frame:
            for para in ph.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.size != TITLE_SIZE:
                        run.font.size = TITLE_SIZE
                        title_fixed += 1

    if i in TWO_CONTENT_SLIDES:
        for ph in slide.placeholders:
            if ph.placeholder_format.idx not in (1, 2) or not ph.has_text_frame:
                continue
            for para in ph.text_frame.paragraphs:
                target = HEADER_SIZE if para.level == 0 else BULLET_SIZE
                for run in para.runs:
                    if run.font.size != target:
                        run.font.size = target
                        body_fixed += 1

prs.save(DECK)
print(f"Titles normalized to 28pt: {title_fixed} runs changed")
print(f"Two-Content body normalized to 18/16pt: {body_fixed} runs changed")
