"""
Append "Team Contributions" to the outline slide's numbered list, so it lists
every distinct topic in the now-16-slide deck rather than the original nine.

Run once, after update_deck_progress.py.

    python docs/patch_outline.py
"""
import os

from pptx import Presentation

HERE = os.path.dirname(os.path.abspath(__file__))
DECK = os.path.join(HERE, "Indus11-Review2.pptx")

prs = Presentation(DECK)
s2 = prs.slides[1]

for sh in s2.shapes:
    if sh.has_text_frame and sh.text_frame.text.startswith("1.  Mentor Feedback"):
        tf = sh.text_frame
        last = tf.paragraphs[-1]
        assert last.text.strip() == "9.  References", last.text
        new_p = tf.add_paragraph()
        template_run = last.runs[0]
        run = new_p.add_run()
        run.text = "10.  Team Contributions"
        run.font.size = template_run.font.size
        run.font.name = template_run.font.name
        run.font.bold = template_run.font.bold
        if template_run.font.color and template_run.font.color.type is not None:
            run.font.color.rgb = template_run.font.color.rgb
        break
else:
    raise SystemExit("outline placeholder not found")

prs.save(DECK)
print("Outline updated")
