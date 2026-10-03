"""
One page to hold during the final review: ~/Downloads/Indus11_CribSheet.docx

Everything a member must be able to say without looking it up, on a single
printable side. Figures are read from docs/eval-results.json, the same file the
report reads, so the sheet cannot quote a number the report does not.

    python docs/build_cribsheet.py
"""
import json
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.expanduser("~/Downloads/Indus11_CribSheet.docx")

with open(os.path.join(HERE, "eval-results.json")) as f:
    EVAL = json.load(f)
CONF = EVAL["metrics"]["confusion"]
FLAG = EVAL["metrics"]["flagged"]
REAL = EVAL["metrics"]["realistic"]
COUNTS = EVAL["counts"]
PADJ = REAL["precision"]
FALSE_PER_CATCH = round((1 - PADJ) / PADJ)

FONT = "Times New Roman"
INK = RGBColor(0, 0, 0)
GREY = RGBColor(0x44, 0x44, 0x44)

doc = Document()
for _s in doc.sections:
    _s.left_margin = _s.right_margin = Inches(0.5)
    _s.top_margin = _s.bottom_margin = Inches(0.45)


def line(text, size=9, bold=False, italic=False, after=1, before=0,
         align=WD_ALIGN_PARAGRAPH.LEFT, colour=INK, indent=0.0):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.0
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.size, r.font.name, r.font.color.rgb = Pt(size), FONT, colour
    return p


def head(text):
    line(text.upper(), 9.5, bold=True, after=2, before=5)


def kv(label, value, size=8.5):
    """A bold label and its value running on one paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(f"{label}  ")
    r.bold = True
    r.font.size, r.font.name, r.font.color.rgb = Pt(size), FONT, INK
    r = p.add_run(value)
    r.font.size, r.font.name, r.font.color.rgb = Pt(size), FONT, INK
    return p


line("INDUS11 — FINAL REVIEW CRIB SHEET", 12, bold=True, after=1,
     align=WD_ALIGN_PARAGRAPH.CENTER)
line("AI-powered financial fraud detection · PRJ_CE_5_2026_7 · 3 October 2026",
     8, italic=True, after=3, align=WD_ALIGN_PARAGRAPH.CENTER, colour=GREY)

head("The one-line answer")
line("A payment arrives; three independent layers score it in parallel; their points add to "
     "0–100; that number decides APPROVE, REVIEW or BLOCK, with a written reason attached. "
     "Rules and graph decide inside the latency budget; the language model explains "
     "afterwards.", 9, after=1)

head("One transaction, start to finish")
kv("1 Gateway", "FastAPI validates the request, loads both account profiles (Redis, falling "
                "back to MongoDB), rate-limits. Scores nothing.")
kv("2 Rule engine 0–40", "blacklist · velocity >5 in 600 s · amount >3× average · merchant · "
                         "risk tier · 25 banking anomaly rules from FIU-IND, RBI, FATF")
kv("3 Graph analyzer 0–30", "Neo4j/Cypher: shared device · shared IP · circular flow (2–4 "
                            "hops, 72 h, time order) · mule chain · fraud proximity")
line("Layers 2 and 3 run together under asyncio.gather — the wait is the slower of the two, "
     "not their sum. The response returns here.", 8, italic=True, after=1, indent=0.15)
kv("4 RAG layer 0–30", "AFTER the response, as a background task: 3 nearest of 58 pattern "
                       "documents from ChromaDB → llama3 via Ollama (temperature 0) → score "
                       "and explanation. Takes 13–14 s.")
kv("5 Decision engine", "S = min(R + G + L, 100). Under 40 APPROVE · 40–69 REVIEW · 70+ "
                        "BLOCK. The dashboard polls every 3 s until the record updates.")

head("The example to memorise — ₹7,04,000 wire into a device and IP cluster")
for row in [
    "Rule engine      high-risk merchant 8 + elevated tier 5 + amount anomaly 12    = 25",
    "Graph layer      shared device 15 + circular flow 12 + proximity 10 = 37, capped = 30",
    "Language model                                                                = 28",
    "Composite        min(25 + 30 + 28, 100) = 83          83 ≥ 70  →  BLOCK",
]:
    line(row, 8, after=0, indent=0.15)
line("Say this: rules and graph alone give 55, a REVIEW. The model's 28 points carry it to "
     "83. No single layer decided it.", 9, bold=True, after=1, before=2)

head("Numbers you must not fumble")
kv("Dataset", f"{COUNTS['total']} labelled — {COUNTS['fraud']} fraud, "
              f"{COUNTS['legit']} legitimate")
kv("Headline", f"precision {FLAG['precision']:.3f} · recall {FLAG['recall']:.3f} · "
               f"F1 {FLAG['f1']:.3f} · false-positive rate "
               f"{REAL['false_positive_rate']:.4f}")
kv("Confusion", f"APPROVE {CONF['APPROVE']['fraud']} fraud / {CONF['APPROVE']['legit']} legit · "
                f"REVIEW {CONF['REVIEW']['fraud']} / {CONF['REVIEW']['legit']} · "
                f"BLOCK {CONF['BLOCK']['fraud']} / {CONF['BLOCK']['legit']}")
kv("Honest precision", f"{PADJ:.3f} at one fraud in a thousand — about {FALSE_PER_CATCH} "
                       f"false alarms per fraud caught (Equation 3.8)")
kv("Latency", "mean 124 ms · median 114 ms · p95 189 ms · max 287 ms · all within 500 ms")
kv("Ablation", "rule alone F1 0.172 · graph alone 0 · model alone 0 · all three 0.981")
kv("Graph layer", "36 of 36 planted ring transactions flagged")
kv("Anomalies", "50 catalogued · 35 detected (30 rule, 5 graph) · 15 not, each listed with "
                "the data it would need")
kv("Engineering", "113 automated tests green in continuous integration · 18 REST routes · "
                  "58 knowledge-base documents")

head("Why it is built this way")
kv("40 / 30 / 30 and 70", "40 < 70, so no single layer can block a customer alone. A "
                          "blacklisted account scores exactly 40 — enough to involve a "
                          "person, not enough to refuse.")
kv("Explanation guard", "rejects any explanation that names no flag which actually fired. "
                        "The one layer that cannot be inspected is the one that gets checked.")
kv("Failed layer", "contributes no points, and an APPROVE is raised to REVIEW. Missing "
                   "evidence is not evidence of safety. A BLOCK already reached is kept.")
kv("Thresholds", "40 and 70 are settings. A sweep suggested 45 and 50 but was derived from "
                 "the same 208 rows, so it was rejected — tuning on the data you then report "
                 "on is how a number stops meaning anything.")

head("If they push")
kv("96 per cent or 7 — which is real?",
   f"Both. {FLAG['precision']:.3f} on a set that is 25 per cent fraud; {PADJ:.3f} at one in "
   f"a thousand. The detector did not change, the mix did. Recall and the false-positive "
   f"rate do not move with prevalence; precision does.")
kv("Synthetic data — does it mean anything?",
   "It shows the pipeline detects what it was built to detect and that no layer does it "
   "alone. It does not establish real-world accuracy, which is why PaySim or IEEE-CIS is "
   "the first item of future work.")
kv("Recall 1.000 — is it perfect?",
   f"No. It missed no fraud in this set. {COUNTS['legit']} legitimate rows are too few to "
   f"resolve a false-alarm rate that matters at one in a thousand, and the cycle check "
   f"still misses the transfer that closes a ring.")
kv("What does the language model add?",
   "Alone it scores zero — capped at 30, it cannot reach 40 by itself. It adds the "
   "explanation, and the points that carry borderline transactions over the line.")

line("Lead with the limitations before you are asked. Stating them first reads as command of "
     "the work, not as a gap in it.", 9, bold=True, after=0, before=5,
     align=WD_ALIGN_PARAGRAPH.CENTER)

doc.save(OUT)
print(f"Saved {OUT}")
